"""Import the supplied DOCX questions and answer tables; emit apply_patch only.

Question numbering and answer letters are source-controlled. Chapter assignment
is curated per question, never inferred from the source's mixed-system groups.
"""
import argparse
import difflib
import json
import re
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SURGERY = {
    1: [31,31,31,31,31,35,35,35,32,1,1,1,3,3,3,3,3],
    2: [4,4,4,4,4,4,4,4,2,6,6,6,5,6,5],
    3: [7,7,7,6,6,9,8,12,12,12,11,13,13,13,13,14,14],
    4: [16,16,16,16,'pancreatitis','pancreatitis',16,16,16,16,16,16,'pancreatitis','pancreatitis','pancreatitis','pancreatitis'],
    5: [17,17,17,19,18,18,20,29,29,28,28,28,28,28,28],
    6: [22,22,26,21,21,21,21,27,27,27,27,23,23,23,23,23,24,24,25,25,25,25,25],
}
BIO = {1:8,2:13,3:12,4:1,5:3,6:6,7:6,10:7,11:7,12:9,14:16,15:20,16:17,17:17,19:18,20:20,21:19,23:14,25:10,27:14,28:14,29:8,30:14}
BIO_TOPICS = {1:'糖代谢',3:'糖代谢',6:'脂代谢',7:'氨基酸与蛋白质',8:'氨基酸与蛋白质',9:'核苷酸代谢',10:'胆色素代谢与生物转化',12:'酶',13:'维生素',14:'小基因',16:'核酸',17:'核酸',18:'氨基酸与蛋白质',19:'核酸',20:'核酸'}

def surgery_topic(n):
    if n == 'pancreatitis' or n in [13,14,15,16]: return '肝胆胰疾病'
    if n >= 30: return '外科总论'
    if n >= 20: return '骨科'
    return {1:'颈部疾病',2:'胸部疾病',3:'乳房疾病',4:'胸部疾病',5:'胃十二指肠疾病',6:'腹部损伤与感染',7:'腹部损伤与感染',8:'小肠与阑尾疾病',9:'小肠与阑尾疾病',10:'结直肠与肛管疾病',11:'结直肠与肛管疾病',12:'腹外疝',17:'周围血管疾病',18:'泌尿外科',19:'泌尿外科'}[n]

def read_doc(path, subject):
    doc = Document(path)
    section = 1
    records = {}
    current = None
    for index,p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text == '参考答案': break
        header = re.match(r'^第(\d+)组（共(\d+)题）$', text)
        if header:
            section = int(header[1]); current = None; continue
        question = re.match(r'^(\d+)/(\d+)\s+(.+)$', text)
        if question:
            number = int(question[1]); assert (section,number) not in records
            current = {'section':section,'number':number,'total':int(question[2]),'text':question[3],'options':{},'sourceParagraph':index+1}
            records[section,number] = current; continue
        option = re.match(r'^([A-D])\.\s*(.+)$', text)
        if option:
            assert current is not None
            assert option[1] not in current['options']
            current['options'][option[1]] = option[2]
        elif text and current:
            raise AssertionError(('Unexpected content',index,text))
    answers = {}
    assert len(doc.tables) == (6 if subject == 'surgery' else 1)
    for section,table in enumerate(doc.tables,1):
        assert [c.text for c in table.rows[0].cells] == ['题号','答案']
        for row in table.rows[1:]:
            q,answer = [c.text.strip() for c in row.cells]
            number,total = map(int,q.split('/'))
            assert (section,number) not in answers
            assert re.fullmatch(r'[A-D]+',answer)
            answers[section,number] = list(answer)
            assert records[section,number]['total'] == total
    assert set(records) == set(answers)
    for identity,q in records.items():
        assert set(q['options']) == set('ABCD'),identity
        q['answer'] = answers[identity]
    if subject == 'surgery':
        assert len(records) == 103
        for s,mapping in SURGERY.items(): assert sorted(n for section,n in records if section==s) == list(range(1,len(mapping)+1))
    else:
        assert {n for _,n in records} == set(BIO)
    return records

def build(path, subject):
    records = read_doc(path,subject)
    groups = []
    used = set()
    for identity,q in records.items():
        if identity in used: continue
        block = [q]
        number = q['number']+1
        while (q['section'],number) in records and records[q['section'],number]['text'].startswith('（接上题）'):
            block.append(records[q['section'],number]); number += 1
        used.update((item['section'],item['number']) for item in block)
        location = SURGERY[q['section']][q['number']-1] if subject=='surgery' else BIO[q['number']]
        lid = 'med-lecture-23' if location=='pancreatitis' else f'lecture-{location:02d}'
        for item in block:
            other = SURGERY[item['section']][item['number']-1] if subject=='surgery' else BIO[item['number']]
            assert location == other, ('case spans chapters',identity)
        topic = surgery_topic(location) if subject=='surgery' else BIO_TOPICS[location]
        suffix = f'{q["number"]:02d}' if len(block)==1 else f'{q["number"]:02d}-{block[-1]["number"]:02d}'
        gid = f'{subject}-teacher-{q["section"]:02d}-{suffix}'
        common = ''
        if len(block)>1:
            case,sep,ask = q['text'].rpartition('。')
            assert sep and ask,identity
            common = case+sep
        options,stems = [],[]
        for item in block:
            category = f'第 {item["number"]} 题选项'
            scope = len(block)>1
            for key,label in item['options'].items():
                option = {'key':f'q{item["number"]}-{key}' if scope else key,'label':label}
                if scope: option.update(displayKey=key,category=category)
                options.append(option)
            text = item['text'].removeprefix('（接上题）').strip()
            if scope and item is q: text = ask
            s = {'number':item['number'],'text':text,'sourceQuestion':item['number'],'sourceParagraph':item['sourceParagraph'],
                 'lectureId':lid,'answer':[f'q{item["number"]}-{k}' if scope else k for k in item['answer']],
                 'answerRaw':''.join(item['answer']),'answerMode':'多选' if len(item['answer'])>1 else '单选','answerState':'文末参考答案'}
            if scope:s['optionCategory']=category
            stems.append(s)
        title = q['text'] if len(block)==1 else f'第 {q["number"]}—{block[-1]["number"]} 题共用病例'
        if len(title)>42:title=title[:41]+'…'
        g = {'id':gid,'page':0,'title':title,'kind':'A','kindLabel':'多项选择' if all(len(s['answer'])>1 for s in stems) else '单项选择',
             'topic':topic,'lectureIds':[lid],'options':options,'stems':stems,'sourceName':path.name,
             'sourceKey':f'{subject}-homework-2026-09','sourceSection':f'{"外科" if subject=="surgery" else "生化"}课后巩固·第{q["section"]}组',
             'sourceQuestion':q['number'] if len(block)==1 else f'{q["number"]}—{block[-1]["number"]}',
             'sourceText':' '.join(item['text'] for item in block),'supplement':True,'hideSource':True,
             'supplementNotice':'已录入文末参考答案','answerSourceLabel':'课后题答案','answerSourceName':path.name}
        if len(block)>1:g.update(sharedStem=common,sharedQuestionCount=len(block))
        groups.append(g)
    lectures=[]
    if subject=='surgery':
        l = next(l for l in json.loads((ROOT/'src/data/med-data.json').read_text())['lectures'] if l['id']=='lecture-23')
        lectures=[{'id':'med-lecture-23','number':23,'title':'内科与外科 胰腺炎（共用第23讲）','pageCount':l['pageCount'],'file':l['file'],'sourceSubject':'med'}]
    return {'meta':{'title':f'{"外科" if subject=="surgery" else "生化"}课后巩固','sourceName':path.name,'groupCount':len(groups),'stemCount':len(records),
                    'answerNote':'按文末参考答案录入，未冒称已按讲义逐项勘误',
                    'missingSourceQuestions':[] if subject=='surgery' else sorted(set(range(1,31))-set(BIO))},'pages':[],'lectures':lectures,'groups':groups}

def main():
    p=argparse.ArgumentParser();p.add_argument('surgery',type=Path);p.add_argument('biochemistry',type=Path);p.add_argument('--check',action='store_true');p.add_argument('--verify',action='store_true');args=p.parse_args()
    results={subject:build(path,subject) for subject,path in [('surgery',args.surgery),('biochemistry',args.biochemistry)]}
    if args.verify:
        for subject,d in results.items():
            assert json.loads((ROOT/f'src/data/{subject}-teacher-supplement.json').read_text()) == d, subject
            print(subject, 'all source stems, scoped options, reference answers and chapter mappings match')
        return
    if args.check:
        for subject,d in results.items():print(subject,json.dumps(d['meta'],ensure_ascii=False))
        return
    print('*** Begin Patch')
    for subject,d in results.items():
        target=ROOT/f'src/data/{subject}-teacher-supplement.json'
        text = json.dumps(d,ensure_ascii=False,indent=2)+'\n'
        if target.exists():
            diff = list(difflib.unified_diff(target.read_text().splitlines(),text.splitlines(),n=3))
            if diff:
                print('*** Update File: '+str(target))
                print('\n'.join('@@' if line.startswith('@@') else line for line in diff[2:]))
        else:
            print('*** Add File: '+str(target))
            print('\n'.join('+'+line for line in text.splitlines()))
    print('*** End Patch')

if __name__=='__main__': main()
